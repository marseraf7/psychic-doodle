#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Công cụ tự kiểm tra luận văn / bài báo tiếng Nga TRƯỚC khi nộp Antiplagiat.

3 chức năng:
  trung-lap     So bản thảo với các tài liệu nguồn bạn đã đọc -> báo cáo HTML tô màu
                từng câu trùng: thiếu trích dẫn / có dẫn nguồn nhưng chép gần nguyên
                văn / trích dẫn đúng. Kèm dò ký tự lạ (chữ Latin lẫn trong chữ Nga,
                ký tự ẩn, chữ trắng) mà Antiplagiat coi là "tài liệu đáng ngờ".
  kiem-tra-dan  Đối chiếu các tham chiếu [N] trong bài với "Список литературы":
                số không tồn tại, tài liệu không được dẫn, trích nguyên văn «...»
                thiếu nguồn hoặc thiếu số trang.
  dinh-dang     Tạo danh mục tài liệu tham khảo theo ГОСТ Р 7.0.100-2018 từ file CSV.

Cách chạy:
    python kiem_tra_luan_van.py trung-lap --ban-thao ban_thao.docx --nguon nguon
    python kiem_tra_luan_van.py kiem-tra-dan --ban-thao ban_thao.docx
    python kiem_tra_luan_van.py dinh-dang --csv tai_lieu.csv
Đọc được .docx, .txt, .pdf (.pdf cần thư viện pypdf).
"""
import argparse, csv, html, os, re, sys, unicodedata
from collections import Counter, defaultdict

# ======================================================================
# Đọc file
# ======================================================================
def _doc_txt(path):
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "cp1251"):   # file tiếng Nga cũ trên Windows hay là cp1251
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")

def _docx_blocks(path):
    """Các đoạn của .docx THEO ĐÚNG THỨ TỰ (gồm cả chữ trong bảng)."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, d)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                for cell in row.cells:
                    yield from cell.paragraphs

# Dấu chú thích chân trang/cuối bài chèn vào chữ của đoạn: '⟨3⟩' (footnote 3), '⟨c3⟩' (endnote 3).
# Dùng ngoặc ⟨⟩ để không lẫn với tham chiếu [N].
_RE_DAU_CT = re.compile(r"⟨c?\d+⟩")
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def _chu_doan_docx(p_el):
    """Chữ của 1 đoạn .docx, có chèn dấu ⟨N⟩ tại vị trí tham chiếu chú thích
    (p.text của python-docx bỏ mất các tham chiếu này)."""
    phan = []
    for el in p_el.iter():
        if el.getparent() is None or el.getparent().tag != _W + "r":
            continue            # bỏ w:tab trong định nghĩa tab-stop của đoạn...
        if el.tag == _W + "t":
            phan.append(el.text or "")
        elif el.tag == _W + "tab":
            phan.append("\t")
        elif el.tag in (_W + "br", _W + "cr"):
            phan.append(" ")
        elif el.tag == _W + "footnoteReference":
            phan.append(f"⟨{el.get(_W + 'id')}⟩")
        elif el.tag == _W + "endnoteReference":
            phan.append(f"⟨c{el.get(_W + 'id')}⟩")
    return "".join(phan)

def doc_chu_thich(path):
    """Chú thích của .docx -> {'⟨3⟩': 'Выготский Л. С. Мышление и речь. М., 1999. С. 15.', ...}."""
    if not path.lower().endswith(".docx"):
        return {}
    import zipfile
    import xml.etree.ElementTree as ET
    kq = {}
    with zipfile.ZipFile(path) as z:
        for ten_file, the, tien_to in (("word/footnotes.xml", "footnote", ""),
                                       ("word/endnotes.xml", "endnote", "c")):
            if ten_file not in z.namelist():
                continue
            goc = ET.fromstring(z.read(ten_file))
            for fn in goc.findall(_W + the):
                if fn.get(_W + "type") in ("separator", "continuationSeparator", "continuationNotice"):
                    continue
                chu = " ".join("".join(t.text or "" for t in para.iter(_W + "t"))
                               for para in fn.iter(_W + "p"))
                kq[f"⟨{tien_to}{fn.get(_W + 'id')}⟩"] = chu.strip()
    return kq

def _noi_dong_txt(dong):
    """File .txt bị ngắt dòng cứng (mỗi dòng ~70 ký tự) -> nối lại thành đoạn.
    Nối dòng với dòng sau khi dòng này chưa hết câu VÀ (dòng sau viết thường, hoặc
    dòng này dài như dòng bị ngắt). Tiêu đề ngắn và mục danh mục (kết thúc bằng '.')
    vẫn đứng riêng."""
    doan, dang = [], ""
    for d in dong:
        d = d.rstrip()
        if not d.strip():
            if dang:
                doan.append(dang)
            doan.append("")
            dang = ""
            continue
        if dang:
            chua_het = not re.search(r"[.!?:;…»\"”)\]]$", dang)
            if dang.endswith("-") and d.lstrip()[:1].islower():
                dang = dang[:-1] + d.lstrip()            # từ bị gạch nối ngắt dòng
                continue
            if chua_het and (d.lstrip()[:1].islower() or len(dang) >= 60):
                dang += " " + d.lstrip()
                continue
            doan.append(dang)
        dang = d
    if dang:
        doan.append(dang)
    return doan

def _nap_pypdf():
    try:
        from pypdf import PdfReader
        return PdfReader
    except BaseException as e:     # pypdf hỏng thư viện phụ có thể ném PanicException (không phải Exception)
        if isinstance(e, (KeyboardInterrupt, SystemExit)):
            raise
        raise RuntimeError("không nạp được thư viện pypdf (cài lại:  pip install --upgrade pypdf "
                           f"cryptography cffi) — {type(e).__name__}: {e}") from None

def doc_doan_van(path):
    """Danh sách đoạn văn (chuỗi) của file."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return [_chu_doan_docx(p._p) for p in _docx_blocks(path)]
    if ext == ".pdf":
        PdfReader = _nap_pypdf()
        text = "\n".join((pg.extract_text() or "") for pg in PdfReader(path).pages)
        # PDF ngắt dòng giữa câu -> nối lại, giữ dòng trống làm ranh giới đoạn
        text = re.sub(r"-\n(?=[а-яёa-z])", "", text)
        return [re.sub(r"\s*\n\s*", " ", p) for p in re.split(r"\n\s*\n", text)]
    if ext in (".txt", ".md", ""):
        return _noi_dong_txt(_doc_txt(path).splitlines())
    raise SystemExit(f"[LỖI] Không đọc được định dạng {ext}: {path} (dùng .docx/.txt/.pdf)")

DUOI_HO_TRO = (".docx", ".txt", ".pdf", ".md")

# ======================================================================
# Xử lý văn bản tiếng Nga
# ======================================================================
TU_DUNG = set("""
и в во не что он на я с со как а то все всё она так его но да ты к у же вы за бы по только
ее её мне было вот от меня еще ещё нет о об из ему когда даже ну ли если уже или ни быть был
него до вас нибудь уж вам ведь там потом себя ей может они тут где есть надо ней для мы
тебя их чем была сам чтоб без чего раз тоже себе под будет ж тогда кто этот того потому
этого какой ним здесь этом один мой тем чтобы нее неё были куда всех можно при два другой
после над больше тот через эти нас про всего них какая много три эту моя свою этой перед
том такой им более всегда всю между это также является являются являлся этих которые который
которая которое которых котором которой такие таких данный данной данного свой своей своих
их лишь именно однако поэтому т е д др см с
the a an of and or to in on at is are was were be been by for with as that this it from
""".split())

_OKONCHANIYA = sorted("""
иями ями ами ого его ому ему ыми ими ых их ой ей ый ий ая яя ое ее ую юю ов ев ам ям ах ях ом ем
ью ия ие ии ию ть ет ют ит ат ят ешь ишь ем им ете ите ала ила ыла ело ило ал ил ыл ла ли ло
а я о е ы и у ю ь й
""".split(), key=len, reverse=True)

_RE_TU = re.compile(r"[A-Za-zА-Яа-яЁё]+(?:-[A-Za-zА-Яа-яЁё]+)*")

def goc_tu(tu):
    """Gốc từ đơn giản (bỏ đuôi biến cách) -> 'психологии'/'психология' về cùng gốc,
    để câu chỉ đổi cách/giống/số vẫn bị nhận ra là trùng."""
    t = tu.lower().replace("ё", "е")
    for dt in ("ся", "сь"):
        if t.endswith(dt) and len(t) - 2 >= 4:
            t = t[:-2]
            break
    for dt in _OKONCHANIYA:
        if t.endswith(dt) and len(t) - len(dt) >= 4:
            return t[:-len(dt)]
    return t

def tu_noi_dung(cau):
    """Gốc các từ có nghĩa trong câu (bỏ từ dừng, số, trích dẫn [..])."""
    cau = _RE_DAU_CT.sub(" ", re.sub(r"\[[^\]]*\]", " ", cau))
    return [goc_tu(w) for w in _RE_TU.findall(cau) if w.lower().replace("ё", "е") not in TU_DUNG]

# Từ viết tắt không kết thúc câu: 'т. е.', 'и т. д.', 'с. 12', 'Л. С. Выготский'...
_VIET_TAT = set("т е д п др пр см рис табл гл им напр ср стр с т.е т.д т.п т.к г гг в вв ок "
                "проф акад канд докт psych et al vol pp p".split())

# Viết tắt thường ĐỨNG CUỐI câu: 'и др.', 'и пр.', 'и т. д.', 'и т. п.', 'в 1990-х гг.'
_RE_VT_CUOI_CAU = re.compile(r"(\bи\s+(др|пр)|\bт\.\s*[дп]|\d\S*\s*гг?|\bвв)\.$", re.I)
_DONG = r"(?:[»\"”)]|⟨c?\d+⟩)*"        # ngoặc đóng / dấu chú thích sau dấu chấm câu

def tach_cau(doan):
    """Tách đoạn thành câu, không cắt sau chữ viết tắt hay chữ cái đầu tên."""
    cau, dau = [], 0
    for m in re.finditer(r"[.!?…]+" + _DONG + r"\s+(?=[«\"„(\[]?[А-ЯЁA-Z0-9])", doan):
        truoc = re.search(r"([A-Za-zА-Яа-яЁё.]+)[.!?…]+" + _DONG + r"\s+$", doan[dau:m.end()])
        tu = truoc.group(1).lower().strip(".") if truoc else ""
        ket_vt = _RE_VT_CUOI_CAU.search(doan[dau:m.start() + 1])
        if doan[m.start()] == "." and not ket_vt and (tu in _VIET_TAT or
                                       (len(tu) == 1 and tu.isalpha())):   # chữ cái đầu tên
            continue
        cau.append(doan[dau:m.end()].strip())
        dau = m.end()
    if doan[dau:].strip():
        cau.append(doan[dau:].strip())
    return cau

# ======================================================================
# Tách phần "Список литературы" khỏi thân bài
# ======================================================================
_RE_TIEU_DE_DM = re.compile(
    r"^\s*(список\s+(использованн\w+\s+)?(литературы|источников(\s+и\s+литературы)?)"
    r"|библиографическ\w+\s+список|литература|references|bibliography)\s*:?\s*$", re.I)
_RE_TIEU_DE_SAU = re.compile(r"^\s*(приложени[ея]\b|appendix)", re.I)

def tach_danh_muc(doan):
    """-> (các đoạn thân bài, các mục trong danh mục tài liệu).
    Lấy tiêu đề danh mục CUỐI CÙNG (dòng trong mục lục thường kèm số trang nên không khớp)."""
    vt = None
    for i, d in enumerate(doan):
        if _RE_TIEU_DE_DM.match(d):
            vt = i
    if vt is None:
        return doan, []
    muc = []
    for d in doan[vt + 1:]:
        if _RE_TIEU_DE_SAU.match(d):
            break
        if d.strip():
            muc.append(d.strip())
    return doan[:vt], muc

# ======================================================================
# Nhận diện trích dẫn trong câu
# ======================================================================
_RE_DAN_SO = re.compile(r"\[\s*\d[^\]]*\]")                               # [12], [5, с. 34]
_RE_DAN_TEN = re.compile(r"\(\s*[А-ЯЁA-Z][^()]{1,80}?,?\s+(1[89]|20)\d{2}[^()]{0,30}\)")  # (Иванов, 2010)
_RE_NGOAC_KEP = re.compile(r"«[^«»]+»|„[^„“]+“|\"[^\"]+\"|“[^“”]+”")

# Выготский (1934) / Иванова и Петров (2010а)
_RE_TEN_NAM = re.compile(r"\b[А-ЯЁA-Z][а-яёa-z-]+(?:\s+(?:и|и\s+др\.|et\s+al\.))?\s*\((1[89]|20)\d{2}[а-яa-z]?\)")

def co_dan_nguon(cau):
    """Có tham chiếu: [N], (Иванов, 2010), Иванов (2010) hoặc chú thích chân trang ⟨N⟩."""
    return bool(_RE_DAN_SO.search(cau) or _RE_DAN_TEN.search(cau)
                or _RE_TEN_NAM.search(cau) or _RE_DAU_CT.search(cau))

def ty_le_ngoac_kep(cau):
    """Tỷ lệ ký tự của câu nằm trong ngoặc kép."""
    trong = sum(len(m.group()) for m in _RE_NGOAC_KEP.finditer(cau))
    return trong / max(1, len(cau.strip()))

def cau_voi_vi_tri(doan):
    """[(câu, vị trí đầu, vị trí cuối trong đoạn), ...]"""
    ds, vt = [], 0
    for c in tach_cau(doan):
        i = doan.find(c, vt)
        i = vt if i < 0 else i
        ds.append((c, i, i + len(c)))
        vt = i + len(c)
    return ds

def phan_tich_doan(doan):
    """Tách đoạn thành câu, kèm cho mỗi câu:
      - 'dan': có trích dẫn — của chính câu, HOẶC của một câu phía sau trong cùng đoạn
        (kiểu viết phổ biến: 1 tham chiếu [N] ở cuối đoạn cho cả đoạn);
      - 'ngoac': tỷ lệ chữ nằm trong ngoặc kép, tính trên cả đoạn để trích dẫn
        «...» kéo dài qua nhiều câu vẫn được nhận ra."""
    ds = cau_voi_vi_tri(doan)
    mask = bytearray(len(doan))
    for m in _RE_NGOAC_KEP.finditer(doan):
        mask[m.start():m.end()] = b"\x01" * (m.end() - m.start())
    kq, co_dan_sau = [], False
    for c, a, b in reversed(ds):
        co_dan_sau = co_dan_sau or co_dan_nguon(c)
        kq.append({"cau": c, "dan": co_dan_sau, "dan_rieng": co_dan_nguon(c),
                   "ngoac": sum(mask[a:b]) / max(1, b - a)})
    return kq[::-1]

# ======================================================================
# Dò ký tự lạ (Antiplagiat gắn cờ "подозрительный документ")
# ======================================================================
_KY_TU_AN = {"​": "zero-width space", "‌": "zero-width non-joiner",
             "‍": "zero-width joiner", "⁠": "word joiner",
             "﻿": "BOM giữa văn bản", "­": "soft hyphen"}
_RE_TU_MOI_CHU = re.compile(r"\w+", re.U)

def _chu_he(c):
    if "а" <= c.lower() <= "я" or c in "ёЁ":
        return "cyr"
    if "a" <= c.lower() <= "z":
        return "lat"
    return None

def do_ky_tu_la(doan):
    """-> danh sách (loại, đoạn trích). Chữ Latin lẫn trong từ Nga thường do chép từ PDF
    hoặc do 'công cụ lách' -> nên sửa lại để bài sạch."""
    loi = []
    for d in doan:
        for kt, ten in _KY_TU_AN.items():
            if kt in d:
                i = d.index(kt)
                loi.append((f"Ký tự ẩn ({ten}) x{d.count(kt)}", d[max(0, i - 40):i + 40].replace(kt, "⟦?⟧")))
        for m in _RE_TU_MOI_CHU.finditer(d):
            w = m.group()
            he = {_chu_he(c) for c in w} - {None}
            if he == {"cyr", "lat"}:
                la = "".join(c for c in w if _chu_he(c) == "lat")
                loi.append((f"Từ lẫn chữ Latin ('{la}') trong chữ Nga", w))
    return loi

def _chuoi_style(st):
    while st is not None:
        yield st
        st = st.base_style

def _mau_rgb(font):
    try:
        if font.color is not None and font.color.type is not None and font.color.rgb is not None:
            return str(font.color.rgb).upper()
    except (AttributeError, ValueError):
        pass
    return None

def _thuoc_tinh_that(r, p):
    """(hidden, cỡ pt, màu RGB) THỰC TẾ của 1 run: lấy giá trị đặt trực tiếp trên run, nếu
    không có thì theo style ký tự của run, rồi style đoạn (kể cả style gốc mà nó kế thừa).
    Nhờ vậy chữ trắng/ẩn đặt qua Style vẫn bị phát hiện."""
    fonts = [r.font]
    try:
        fonts += [st.font for st in _chuoi_style(r.style)]
    except (KeyError, ValueError, AttributeError):
        pass
    try:
        fonts += [st.font for st in _chuoi_style(p.style)]
    except (KeyError, ValueError, AttributeError):
        pass
    an = next((f.hidden for f in fonts if f.hidden is not None), None)
    co = next((f.size.pt for f in fonts if f.size is not None), None)
    mau = next((m for m in map(_mau_rgb, fonts) if m is not None), None)
    return an, co, mau

def do_chu_an_docx(path):
    """Chữ màu trắng / cỡ < 3pt / thuộc tính hidden trong .docx (đặt trực tiếp hoặc qua Style)."""
    loi = []
    for p in _docx_blocks(path):
        for r in p.runs:
            if not r.text.strip():
                continue
            an, co, mau = _thuoc_tinh_that(r, p)
            ly_do = None
            if an:
                ly_do = "chữ bị ẩn (hidden)"
            elif co is not None and co < 3:
                ly_do = f"chữ cỡ {co:g}pt"
            elif mau == "FFFFFF":
                ly_do = "chữ màu trắng"
            if ly_do:
                loi.append((ly_do, r.text[:80]))
    return loi

# ======================================================================
# 1) KIỂM TRA TRÙNG LẶP
# ======================================================================
N_SHINGLE = 3       # cụm 3 từ có nghĩa liên tiếp
MIN_TU_CAU = 5      # câu ngắn hơn -> bỏ qua (quá ngắn để kết luận)

def _shingles(goc):
    return [tuple(goc[i:i + N_SHINGLE]) for i in range(len(goc) - N_SHINGLE + 1)]

class ChiMucNguon:
    """Chỉ mục cụm từ của tất cả tài liệu nguồn."""
    def __init__(self):
        self.ten = []            # id nguồn -> tên file
        self.cau = []            # id nguồn -> danh sách câu
        self.index = defaultdict(set)   # shingle -> {(id nguồn, id câu)}

    def them(self, ten, doan):
        sid = len(self.ten)
        self.ten.append(ten)
        cau_list = [c for d in doan for c in tach_cau(d)]
        self.cau.append(cau_list)
        for cid, c in enumerate(cau_list):
            for sh in _shingles(tu_noi_dung(c)):
                self.index[sh].add((sid, cid))

    def so_khop(self, cau):
        """-> (độ phủ 0..1, số từ có nghĩa, id nguồn khớp nhất, câu nguồn khớp nhất)."""
        goc = tu_noi_dung(cau)
        if len(goc) < MIN_TU_CAU:
            return 0.0, len(goc), None, None
        phu = [False] * len(goc)
        dem = Counter()
        for i, sh in enumerate(_shingles(goc)):
            hit = self.index.get(sh)
            if hit:
                for k in range(i, i + N_SHINGLE):
                    phu[k] = True
                for sc in hit:
                    dem[sc] += 1
        if not dem:
            return 0.0, len(goc), None, None
        (sid, cid), _ = dem.most_common(1)[0]
        return sum(phu) / len(goc), len(goc), sid, self.cau[sid][cid]

# Trạng thái mỗi câu
OK, TRICH_DUNG, CAN_DIEN_DAT, THIEU_NGUON = "ok", "trich_dung", "can_dien_dat", "thieu_nguon"
MO_TA = {
    THIEU_NGUON:  ("Trùng nguồn, KHÔNG có trích dẫn",
                   "Thêm tham chiếu [N, с. X]; nếu giữ nguyên văn thì đặt trong «...». "
                   "Tốt nhất: tóm ý chính bằng lời của bạn rồi phân tích/đánh giá thêm."),
    CAN_DIEN_DAT: ("Có dẫn nguồn nhưng chép gần nguyên văn, không ngoặc kép",
                   "Hoặc đặt đoạn này trong «...» kèm số trang (trích dẫn trực tiếp), "
                   "hoặc tóm lược ý tác giả bằng lời của bạn và giữ tham chiếu."),
    TRICH_DUNG:   ("Trích dẫn trực tiếp đúng cách",
                   "Antiplagiat tính vào mục «цитирования», không phải đạo văn. "
                   "Lưu ý: tổng trích dẫn không nên chiếm quá nhiều bài."),
}

def phan_loai(do_phu, nguong, dan, ngoac):
    if do_phu < nguong:
        return OK
    if dan and ngoac >= 0.4:
        return TRICH_DUNG
    return CAN_DIEN_DAT if dan else THIEU_NGUON

NGUON_RONG = 20     # nguồn có ít hơn số từ này -> gần như chắc là PDF scan / file lỗi

def kiem_tra_trung_lap(ban_thao, thu_muc_nguon, nguong=0.5):
    """-> dict kết quả (dùng cho báo cáo và test)."""
    doan = doc_doan_van(ban_thao)
    than_bai, _ = tach_danh_muc(doan)

    chi_muc = ChiMucNguon()
    canh_bao = []
    files = []
    if os.path.isdir(thu_muc_nguon):
        for goc, _, ds in os.walk(thu_muc_nguon):
            files += [os.path.join(goc, f) for f in sorted(ds)
                      if f.lower().endswith(DUOI_HO_TRO) and not f.startswith("~$")]
    elif os.path.isfile(thu_muc_nguon):
        files = [thu_muc_nguon]
    if not files:
        raise SystemExit(f"[LỖI] Không có tài liệu nguồn (.docx/.txt/.pdf) trong: {thu_muc_nguon}")
    for f in files:
        try:
            nguon_doan = doc_doan_van(f)
        except SystemExit:
            raise
        except Exception as e:   # 1 file hỏng không làm dừng cả lượt kiểm tra
            canh_bao.append(f"{os.path.basename(f)}: KHÔNG đọc được, đã bỏ qua — {e}")
            print(f"  [BỎ QUA] {canh_bao[-1]}")
            continue
        ten = os.path.relpath(f, thu_muc_nguon) if os.path.isdir(thu_muc_nguon) else os.path.basename(f)
        so_tu = sum(len(tu_noi_dung(d)) for d in nguon_doan)
        if so_tu < NGUON_RONG:
            canh_bao.append(f"{ten}: gần như không có chữ ({so_tu} từ) — nếu là PDF scan thì cần "
                            "OCR trước, nếu không nguồn này KHÔNG được so sánh.")
            print(f"  [CẢNH BÁO] {canh_bao[-1]}")
        chi_muc.them(ten, nguon_doan)

    ket_qua_doan, tong_tu = [], 0
    tu_theo_loai, tu_theo_nguon = Counter(), Counter()
    for d in than_bai:
        ds_cau = []
        for pt in phan_tich_doan(d):
            c = pt["cau"]
            do_phu, so_tu, sid, cau_nguon = chi_muc.so_khop(c)
            loai = phan_loai(do_phu, nguong, pt["dan"], pt["ngoac"])
            tong_tu += so_tu
            if loai != OK:
                tu_theo_loai[loai] += so_tu
                tu_theo_nguon[chi_muc.ten[sid]] += so_tu
            ds_cau.append({"cau": c, "do_phu": do_phu, "loai": loai,
                           "nguon": chi_muc.ten[sid] if sid is not None else None,
                           "cau_nguon": cau_nguon})
        ket_qua_doan.append(ds_cau)

    ky_tu_la = do_ky_tu_la(doan)
    if ban_thao.lower().endswith(".docx"):
        ky_tu_la += do_chu_an_docx(ban_thao)

    return {"ban_thao": ban_thao, "nguon": chi_muc.ten, "nguong": nguong,
            "doan": ket_qua_doan, "tong_tu": tong_tu,
            "tu_theo_loai": tu_theo_loai, "tu_theo_nguon": tu_theo_nguon,
            "ky_tu_la": ky_tu_la, "canh_bao_nguon": canh_bao}

def _pt(a, b):
    return 100.0 * a / b if b else 0.0

MAU = {THIEU_NGUON: "#f8c9c4", CAN_DIEN_DAT: "#fbe3a8", TRICH_DUNG: "#cdebc5"}
MAU_TOI = {THIEU_NGUON: "#6b2a24", CAN_DIEN_DAT: "#5e4a12", TRICH_DUNG: "#28502a"}

def ghi_bao_cao_html(kq, out):
    e = html.escape
    T = kq["tong_tu"]
    hang_loai = "".join(
        f'<tr><td><span class="o {k}"></span>{e(MO_TA[k][0])}</td>'
        f'<td class="so">{_pt(kq["tu_theo_loai"][k], T):.1f}%</td><td>{e(MO_TA[k][1])}</td></tr>'
        for k in (THIEU_NGUON, CAN_DIEN_DAT, TRICH_DUNG))
    hang_nguon = "".join(
        f'<tr><td>{e(t)}</td><td class="so">{_pt(n, T):.1f}%</td></tr>'
        for t, n in kq["tu_theo_nguon"].most_common()) or '<tr><td colspan="2">Không có câu trùng.</td></tr>'
    than = []
    for ds in kq["doan"]:
        if not ds:
            continue
        phan = []
        for c in ds:
            if c["loai"] == OK:
                phan.append(e(c["cau"]))
            else:
                tip = (f'{MO_TA[c["loai"]][0]} — trùng {c["do_phu"] * 100:.0f}% với «{c["nguon"]}»:\n'
                       f'{c["cau_nguon"]}')
                phan.append(f'<mark class="{c["loai"]}" title="{e(tip)}">{e(c["cau"])}</mark>')
        than.append("<p>" + " ".join(phan) + "</p>")
    la = "".join(f"<li><b>{e(l)}</b>: <code>{e(s)}</code></li>" for l, s in kq["ky_tu_la"][:200])
    cb = "".join(f"<li>{e(x)}</li>" for x in kq.get("canh_bao_nguon", []))
    khoi_la = ""
    if cb:
        khoi_la += ('<section class="canh-bao"><h2>Tài liệu nguồn không được so sánh</h2>'
                    f"<ul>{cb}</ul></section>")
    if kq["ky_tu_la"]:
        khoi_la += (f'<section class="canh-bao"><h2>Ký tự bất thường ({len(kq["ky_tu_la"])})</h2>'
                    "<p>Antiplagiat gắn cờ «подозрительный документ» khi thấy những thứ này, kể cả khi "
                    "do vô tình chép từ PDF. Hãy gõ lại các từ này / xóa ký tự ẩn.</p>"
                    f"<ul>{la}</ul></section>")
    trang = f"""<!doctype html><html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Báo cáo trùng lặp</title><style>
:root{{--nen:#fbfaf7;--chu:#1f1f1f;--phu:#666;--vien:#ddd;
{''.join(f'--{k}:{v};' for k, v in MAU.items())}}}
@media (prefers-color-scheme: dark){{:root{{--nen:#17181a;--chu:#e8e6e1;--phu:#a0a0a0;--vien:#3a3a3a;
{''.join(f'--{k}:{v};' for k, v in MAU_TOI.items())}}}}}
body{{background:var(--nen);color:var(--chu);font:16px/1.6 Georgia,"Times New Roman",serif;
max-width:900px;margin:0 auto;padding:24px 16px}}
h1,h2{{font-family:system-ui,sans-serif}} table{{border-collapse:collapse;width:100%;margin:8px 0 24px}}
td,th{{border-bottom:1px solid var(--vien);padding:6px 8px;vertical-align:top;text-align:left}}
.so{{text-align:right;white-space:nowrap;font-family:system-ui}} .phu{{color:var(--phu);font-size:14px}}
mark{{color:inherit;border-radius:3px;padding:0 2px;cursor:help}}
{''.join(f'mark.{k},.o.{k}{{background:var(--{k})}}' for k in MAU)}
.o{{display:inline-block;width:12px;height:12px;border-radius:2px;margin-right:6px}}
.canh-bao{{border-left:4px solid #c0392b;padding:4px 12px;margin:16px 0}}
code{{word-break:break-all}} .bai p{{text-align:justify}}
</style></head><body>
<h1>Báo cáo tự kiểm tra trùng lặp</h1>
<p class="phu">Bản thảo: {e(os.path.basename(kq["ban_thao"]))} · {len(kq["nguon"])} tài liệu nguồn ·
ngưỡng {kq["nguong"] * 100:.0f}% · {T} từ có nghĩa (không tính danh mục tài liệu)</p>
<p class="phu">Đây là so sánh với <b>các nguồn bạn cung cấp</b>, không thay thế Antiplagiat
(Antiplagiat so với cơ sở dữ liệu rất lớn). Mục đích: biết chỗ nào cần dẫn nguồn đúng,
chỗ nào nên tự viết lại bằng ý của mình.</p>
<h2>Tổng quan</h2><table>{hang_loai}</table>
<h2>Theo tài liệu nguồn</h2><table>{hang_nguon}</table>
{khoi_la}
<h2>Toàn văn</h2><p class="phu">Rê chuột lên đoạn tô màu để xem câu nguồn tương ứng.</p>
<div class="bai">{''.join(than)}</div>
</body></html>"""
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        f.write(trang)

def lenh_trung_lap(a):
    print(f"Đang so sánh {a.ban_thao} với nguồn trong {a.nguon} ...")
    kq = kiem_tra_trung_lap(a.ban_thao, a.nguon, a.nguong)
    T = kq["tong_tu"]
    print(f"  {len(kq['nguon'])} tài liệu nguồn, {T} từ có nghĩa trong thân bài.")
    for k in (THIEU_NGUON, CAN_DIEN_DAT, TRICH_DUNG):
        print(f"  {MO_TA[k][0]:<60} {_pt(kq['tu_theo_loai'][k], T):5.1f}%")
    if kq["ky_tu_la"]:
        print(f"  [CẢNH BÁO] {len(kq['ky_tu_la'])} ký tự/từ bất thường — xem cuối báo cáo.")
    ghi_bao_cao_html(kq, a.out)
    print(f"Đã ghi báo cáo: {a.out}")

# ======================================================================
# 2a) KIỂM TRA TRÍCH DẪN
# ======================================================================
def _so_trong_ngoac(noi_dung):
    """'5, с. 12; 7–9' -> ({5, 7, 8, 9}, co_so_trang)."""
    so, co_trang = set(), False
    for phan in noi_dung.split(";"):
        m = re.search(r"\b(с|c|стр|p|pp|s)\.\s*\d", phan, re.I)
        if m:
            co_trang = True
            phan = phan[:m.start()]
        for a, b in re.findall(r"(\d+)(?:\s*[-–—]\s*(\d+))?", phan):
            a = int(a)
            b = int(b) if b else a
            if 0 < a <= b <= 999 and b - a < 100:   # [2015] là năm, không phải số thứ tự
                so.update(range(a, b + 1))
    return so, co_trang

_RE_NAM = re.compile(r"\b(1[5-9]|20)\d{2}\b")
_RE_CO_TRANG = re.compile(r"(?<![А-ЯЁа-яёA-Za-z])(с|c|стр|p|pp|s)\.\s*\d", re.I)

def _ho_trong_muc(muc):
    """Gốc họ tác giả trong 1 mục danh mục: 'Выготский, Л. С.', 'Иванов И. И.', '/ Л. С. Выготский'."""
    muc = re.sub(r"^\s*\d+[.)]\s*", "", muc)
    ho = re.findall(r"([А-ЯЁA-Z][а-яёa-z'-]+),?\s+[А-ЯЁA-Z]\.", muc)
    ho += re.findall(r"(?:[А-ЯЁA-Z]\.\s?){1,2}([А-ЯЁA-Z][а-яёa-z'-]+)", muc)
    return {goc_tu(h) for h in ho}

def trich_dan_tac_gia_nam(text):
    """Trích dẫn kiểu tác giả–năm -> [(gốc họ, chuỗi gốc)]:
    '(Иванов, 2010; Петров и др., 2011)', 'Выготский (1934)'."""
    kq = []
    for m in re.finditer(r"\(([^()]{2,200})\)", text):
        for phan in m.group(1).split(";"):
            mm = re.match(r"\s*(?:см\.\s*|cf\.\s*)?([А-ЯЁA-Z][а-яёa-z'-]+)", phan)
            if mm and _RE_NAM.search(phan):
                kq.append((goc_tu(mm.group(1)), phan.strip()))
    for m in _RE_TEN_NAM.finditer(text):
        kq.append((goc_tu(m.group().split()[0]), m.group()))
    return kq

def _tham_chieu_sat(text):
    """Tham chiếu đứng NGAY đầu chuỗi (sau dấu câu/khoảng trắng): [..], (Tác giả, năm), ⟨N⟩."""
    text = re.sub(r"^[\s,.;:!?…»\"”]*", "", text)
    for rx in (r"\[[^\]]+\]", _RE_DAU_CT.pattern, _RE_DAN_TEN.pattern):
        m = re.match(rx, text)
        if m:
            return m.group()
    return None

def kiem_tra_trich_dan(ban_thao, min_tu_trich=6):
    doan = doc_doan_van(ban_thao)
    chu_thich = doc_chu_thich(ban_thao)
    than_bai, danh_muc = tach_danh_muc(doan)
    van_de = []          # (mức, nội dung)
    da_dan = Counter()   # số thứ tự [N] -> số lần
    tg_nam = []          # (gốc họ, chuỗi) của trích dẫn tác giả–năm
    so_chu_thich = 0

    def co_trang(tc):
        if tc is None:
            return False
        if tc.startswith("["):
            return _so_trong_ngoac(tc[1:-1])[1]
        if tc.startswith("⟨"):
            return bool(_RE_CO_TRANG.search(chu_thich.get(tc, "")))
        return bool(_RE_CO_TRANG.search(tc))

    for d in than_bai:
        for m in re.finditer(r"\[([^\[\]]{1,120})\]", d):
            for n in _so_trong_ngoac(m.group(1))[0]:
                da_dan[n] += 1
        tg_nam += trich_dan_tac_gia_nam(d)
        so_chu_thich += len(_RE_DAU_CT.findall(d))

        cau = cau_voi_vi_tri(d)
        for q in _RE_NGOAC_KEP.finditer(d):   # cả đoạn: trích dẫn «...» có thể dài nhiều câu
            if len(_RE_TU.findall(q.group())) < min_tu_trich:
                continue   # «...» ngắn thường là thuật ngữ / tên riêng, không phải trích dẫn
            tc = _tham_chieu_sat(d[q.end():q.end() + 80])
            if tc is None:   # tham chiếu đặt TRƯỚC trích dẫn trong cùng câu: 'Выготский (1934) писал: «...»'
                dau_cau = max((a for _, a, _ in cau if a <= q.start()), default=0)
                truoc = d[dau_cau:q.start()]
                tim = [m.group() for rx in (r"\[[^\]]+\]", _RE_DAU_CT.pattern, _RE_DAN_TEN.pattern,
                                            _RE_TEN_NAM.pattern) for m in re.finditer(rx, truoc)]
                if not tim:
                    van_de.append(("LỖI", f"Trích nguyên văn nhưng không có tham chiếu: {q.group()[:160]}"))
                    continue
                tc = tim[-1]
            if not co_trang(tc):
                van_de.append(("NHẮC", f"Trích nguyên văn nên ghi số trang (с. X): {q.group()[:160]}"))

    ho_ct = {goc_tu(w) for t in chu_thich.values() for w in _RE_TU.findall(t)}
    so_muc = len(danh_muc)
    if not (da_dan or tg_nam or so_chu_thich):
        van_de.append(("LỖI", "Không tìm thấy trích dẫn nào: [N], (Tác giả, năm) hay chú thích chân trang."))
    if not danh_muc:
        van_de.append(("LỖI", "Không tìm thấy tiêu đề «Список литературы» / «Список использованных "
                              "источников» (tiêu đề phải nằm riêng một dòng)."))
    else:
        ho_muc = [_ho_trong_muc(m) for m in danh_muc]
        tat_ca_ho = set().union(*ho_muc)
        for n in sorted(da_dan):
            if n > so_muc:
                van_de.append(("LỖI", f"Tham chiếu [{n}] ({da_dan[n]} lần) nhưng danh mục chỉ có {so_muc} mục."))
        bao = set()
        for goc, chuoi in tg_nam:
            if goc not in tat_ca_ho and goc not in bao:
                bao.add(goc)
                van_de.append(("LỖI", f"Trích dẫn «{chuoi}» nhưng không có tác giả này trong danh mục."))
        ho_da_dan = {g for g, _ in tg_nam} | ho_ct
        for i, (muc, ho) in enumerate(zip(danh_muc, ho_muc), 1):
            if i not in da_dan and not (ho & ho_da_dan):
                van_de.append(("NHẮC", f"Mục {i} trong danh mục chưa được dẫn lần nào: {muc[:120]}"))
            for loi in kiem_tra_muc_gost(muc):
                van_de.append(("NHẮC", f"Mục {i}: {loi} — {muc[:120]}"))
    return {"so_muc": so_muc, "da_dan": da_dan, "tg_nam": tg_nam,
            "so_chu_thich": so_chu_thich, "van_de": van_de}

def kiem_tra_muc_gost(muc):
    """Kiểm tra nhanh các thành phần bắt buộc của 1 mục (không thay được người hướng dẫn)."""
    loi = []
    muc = re.sub(r"^\s*\d+[.)]\s*", "", muc)
    if not re.search(r"\b(1[5-9]|20)\d{2}\b", muc):
        loi.append("thiếu năm xuất bản")
    if re.search(r"https?://|www\.", muc):
        if "URL" not in muc:
            loi.append("tài liệu điện tử nên ghi 'URL: ...'")
        if not re.search(r"дата обращения|accessed", muc, re.I):
            loi.append("thiếu '(дата обращения: ДД.ММ.ГГГГ)'")
    elif "//" in muc:
        if not re.search(r"\b[СCP]\.\s*\d", muc):
            loi.append("bài báo thiếu số trang 'С. 45–56'")
    elif not re.search(r"\d+\s*[сcp]\.", muc):
        loi.append("sách thiếu tổng số trang '352 с.'")
    if " - " in muc:
        loi.append("dùng gạch nối '-' thay cho gạch ngang ' – ' giữa các vùng")
    return loi

def lenh_kiem_tra_dan(a):
    kq = kiem_tra_trich_dan(a.ban_thao)
    dong = [f"KIỂM TRA TRÍCH DẪN: {a.ban_thao}",
            f"Danh mục: {kq['so_muc']} mục · [N]: {len(kq['da_dan'])} số khác nhau · "
            f"tác giả–năm: {len(kq['tg_nam'])} · chú thích chân trang: {kq['so_chu_thich']}", ""]
    if not kq["van_de"]:
        dong.append("Không phát hiện vấn đề.")
    for muc, nd in kq["van_de"]:
        dong.append(f"[{muc}] {nd}")
    text = "\n".join(dong)
    print(text)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as f:
        f.write(text + "\n")
    print(f"\nĐã ghi: {a.out}")

# ======================================================================
# 2b) ĐỊNH DẠNG DANH MỤC THEO ГОСТ Р 7.0.100-2018
# ======================================================================
_RE_CHU_TAT = re.compile(r"^(?:[А-ЯЁA-Z]\.)+$")

def tach_ten(s):
    """'Выготский Л. С.' / 'Выготский, Л.С.' / 'Выготский Лев Семенович' / 'Л. С. Выготский'
    -> ('Выготский', 'Л. С.')."""
    s = s.strip().replace(",", " ")
    phan = s.split()
    if not phan:
        return None
    # Chữ tắt viết TRƯỚC họ ('Л. С. Выготский', 'Л.С. Выготский') -> đưa họ lên đầu
    if len(phan) > 1 and _RE_CHU_TAT.match(phan[0]) and not _RE_CHU_TAT.match(phan[-1]):
        phan = [phan[-1]] + phan[:-1]
    ho, du = phan[0], phan[1:]
    tat = []
    for p in du:
        for x in filter(None, p.split(".")):
            tat.append(x[0].upper() + ".")
    return ho, " ".join(tat)

def _tac_gia(s):
    return [t for t in (tach_ten(x) for x in re.split(r";", s or "")) if t]

def _gach(s):
    return re.sub(r"\s*[-—]\s*", "–", s.strip())   # 45-56 -> 45–56

def _cham(s):
    """Kết thúc bằng dấu chấm (không nhân đôi)."""
    s = s.strip()
    return s if s.endswith((".", "?", "!", "…")) else s + "."

LOAI = ("sach", "bai_bao", "luan_an", "tom_tat", "web")
_BI_DANH = {"book": "sach", "kniga": "sach", "книга": "sach", "article": "bai_bao",
            "статья": "bai_bao", "baibao": "bai_bao", "thesis": "luan_an", "dissertation": "luan_an",
            "диссертация": "luan_an", "luanan": "luan_an", "luanvan": "luan_an",
            "автореферат": "tom_tat", "tomtat": "tom_tat", "website": "web", "сайт": "web", "url": "web"}
BAT_BUOC = {"sach": ("ten", "thanh_pho", "nam", "so_trang"),
            "bai_bao": ("ten", "tap_chi", "nam", "trang"),
            "luan_an": ("tac_gia", "ten", "thanh_pho", "nam", "so_trang"),
            "tom_tat": ("tac_gia", "ten", "thanh_pho", "nam", "so_trang"),
            "web": ("ten", "url", "ngay_truy_cap")}

def chuan_loai(x):
    """'Sách' / 'bài báo' / 'book' / 'Статья' -> mã loại chuẩn; không nhận ra -> None."""
    t = "".join(c for c in unicodedata.normalize("NFD", (x or "").strip().lower())
                if unicodedata.category(c) != "Mn").replace("đ", "d")
    if not t:
        return ""
    t2 = re.sub(r"[\s_-]+", "", t)
    for l in LOAI:
        if t2 == l.replace("_", ""):
            return l
    return _BI_DANH.get(t2) or _BI_DANH.get(t)

def kiem_tra_dong(r):
    """-> danh sách lỗi của 1 dòng CSV (rỗng = hợp lệ)."""
    g = lambda k: (r.get(k) or "").strip()
    loai = chuan_loai(g("loai"))
    if loai is None:
        return [f"loại '{g('loai')}' không hợp lệ (dùng: {', '.join(LOAI)})"]
    thieu = [c for c in BAT_BUOC[loai or "sach"] if not g(c)]
    return [f"thiếu cột: {', '.join(thieu)}"] if thieu else []

def dinh_dang_muc(r):
    """1 dòng CSV (dict) -> 1 mục danh mục theo ГОСТ Р 7.0.100-2018."""
    g = lambda k: (r.get(k) or "").strip()
    loai = chuan_loai(g("loai")) or "sach"
    tg = _tac_gia(g("tac_gia"))
    # ≤3 tác giả: tiêu đề mục là tác giả đầu; ≥4: bắt đầu bằng tên tài liệu
    dau = f"{tg[0][0]}, {tg[0][1]} ".replace(",  ", ", ") if 1 <= len(tg) <= 3 else ""
    ds = tg if len(tg) <= 4 else tg[:3]
    trach_nhiem = ", ".join(f"{i} {h}".strip() for h, i in ds) + (" [и др.]" if len(tg) > 4 else "")

    ten = g("ten")
    if g("thong_tin_them"):
        ten += " : " + g("thong_tin_them")
    if loai in ("luan_an", "tom_tat"):
        bac = g("bac") or "кандидата"
        nganh = g("nganh") or "психологических"
        dang = ("автореферат диссертации" if loai == "tom_tat" else "диссертация")
        if g("chuyen_nganh"):
            ten += f" : специальность {g('chuyen_nganh')}"
        ten += f" : {dang} на соискание ученой степени {bac} {nganh} наук"
    tn = " / " + trach_nhiem if trach_nhiem else ""
    if g("bien_tap"):
        tn += (" ; " if tn else " / ") + g("bien_tap")
    if loai in ("luan_an", "tom_tat") and len(tg) == 1 and "." not in g("tac_gia"):
        tn = " / " + re.sub(r"[\s,]+", " ", g("tac_gia"))   # luận án: 'Фамилия Имя Отчество'
    if loai in ("luan_an", "tom_tat") and g("noi_bao_ve"):
        tn += " ; " + g("noi_bao_ve")
    muc = dau + ten + tn

    if loai == "bai_bao":
        muc += f" // {g('tap_chi')}. – {g('nam')}"
        so = ", ".join(x for x in (f"Т. {g('tap')}" if g("tap") else "",
                                   f"№ {g('so')}" if g("so") else "") if x)
        if so:
            muc += f". – {so}"
        if g("trang"):
            muc += f". – С. {_gach(g('trang'))}"
        muc = _cham(muc)
    elif loai == "web":
        muc = _cham(muc) + " – Текст : электронный // " + (g("ten_trang_web") or "Сайт") + " : [сайт]"
        if g("nam"):
            muc += f". – {g('nam')}"
        muc += f". – URL: {g('url')}"
        if g("ngay_truy_cap"):
            muc += f" (дата обращения: {g('ngay_truy_cap')})"
        muc = _cham(muc)
    else:   # sach, luan_an, tom_tat
        muc = _cham(muc) + " – "
        noi = g("thanh_pho")
        if loai == "sach" and g("nha_xb"):
            noi += f" : {g('nha_xb')}"
        muc += f"{noi}, {g('nam')}" if noi else g("nam")
        if g("so_trang"):
            # sách tiếng nước ngoài ghi 'p.', tiếng Nga ghi 'с.'
            muc += f". – {g('so_trang')} " + ("p" if _khoa_sap_xep(ten)[0] else "с")
        if g("isbn"):
            muc += f". – ISBN {g('isbn')}"
        muc = _cham(muc)
    return re.sub(r"\s{2,}", " ", muc).replace(" .", ".")

def _khoa_sap_xep(muc):
    """Tài liệu tiếng Nga trước, tiếng nước ngoài sau; mỗi nhóm theo bảng chữ cái."""
    c = next((ch for ch in muc if ch.isalpha()), "")
    return (0 if _chu_he(c) == "cyr" else 1, muc.lower().replace("ё", "е"))

def lenh_dinh_dang(a):
    import io
    noi_dung = _doc_txt(a.csv)          # Excel tiếng Nga lưu CSV bằng cp1251
    tieu_de = noi_dung.split("\n", 1)[0]
    # Excel bản tiếng Nga/Việt lưu CSV bằng ';' -> chọn dấu xuất hiện nhiều nhất ở dòng tiêu đề
    sep = max(",;\t", key=tieu_de.count)
    muc, so_loi = [], 0
    for i, r in enumerate(csv.DictReader(io.StringIO(noi_dung, newline=""), delimiter=sep), 2):
        if not any((v or "").strip() for v in r.values() if isinstance(v, str)):
            continue
        loi = kiem_tra_dong(r)
        if loi:
            so_loi += 1
            print(f"  [BỎ QUA] Dòng {i} ({(r.get('ten') or '').strip()[:50]}): {'; '.join(loi)}")
            continue
        muc.append(dinh_dang_muc(r))
    if so_loi:
        print(f"  [CẢNH BÁO] {so_loi} dòng bị bỏ qua vì thiếu thông tin — sửa trong CSV rồi chạy lại.")
    if not a.giu_thu_tu:
        muc.sort(key=_khoa_sap_xep)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    txt = os.path.splitext(a.out)[0] + ".txt"
    with open(txt, "w", encoding="utf-8") as f:
        f.write("Список использованных источников\n\n")
        f.write("\n".join(f"{i}. {m}" for i, m in enumerate(muc, 1)) + "\n")
    print(f"Đã tạo {len(muc)} mục -> {txt}")
    if a.out.lower().endswith(".docx"):
        try:
            import docx
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt, Cm
        except ImportError:
            print("  (Không có python-docx nên chỉ ghi .txt:  pip install python-docx)")
            return
        d = docx.Document()
        st = d.styles["Normal"]
        st.font.name, st.font.size = "Times New Roman", Pt(14)
        h = d.add_paragraph("Список использованных источников")
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.runs[0].bold = True
        for i, m in enumerate(muc, 1):
            p = d.add_paragraph(f"{i}. {m}")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(1.25)
        d.save(a.out)
        print(f"Đã tạo -> {a.out}")

# ======================================================================
def main(argv=None):
    for s in (sys.stdout, sys.stderr):   # in tiếng Việt/Nga không lỗi trên console Windows
        try:
            s.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="Tự kiểm tra trùng lặp và trích dẫn cho luận văn tiếng Nga")
    sub = ap.add_subparsers(dest="lenh", required=True)

    p = sub.add_parser("trung-lap", help="So bản thảo với tài liệu nguồn -> báo cáo HTML")
    p.add_argument("--ban-thao", required=True, help="File bản thảo (.docx/.txt/.pdf)")
    p.add_argument("--nguon", required=True, help="Thư mục (hoặc 1 file) tài liệu nguồn")
    p.add_argument("--nguong", type=float, default=0.5,
                   help="Câu có ≥ tỷ lệ này từ trùng thì tô màu (mặc định 0.5)")
    p.add_argument("--out", default=os.path.join("ket_qua", "BAO_CAO_TRUNG_LAP.html"))
    p.set_defaults(ham=lenh_trung_lap)

    p = sub.add_parser("kiem-tra-dan", help="Đối chiếu [N] trong bài với Список литературы")
    p.add_argument("--ban-thao", required=True)
    p.add_argument("--out", default=os.path.join("ket_qua", "BAO_CAO_TRICH_DAN.txt"))
    p.set_defaults(ham=lenh_kiem_tra_dan)

    p = sub.add_parser("dinh-dang", help="CSV -> danh mục theo ГОСТ Р 7.0.100-2018")
    p.add_argument("--csv", required=True)
    p.add_argument("--out", default=os.path.join("ket_qua", "DANH_MUC_TAI_LIEU.docx"))
    p.add_argument("--giu-thu-tu", action="store_true",
                   help="Giữ thứ tự như trong CSV (mặc định: sắp xếp chữ cái, tiếng Nga trước)")
    p.set_defaults(ham=lenh_dinh_dang)

    a = ap.parse_args(argv)
    try:
        a.ham(a)
    except RuntimeError as e:        # vd. không đọc được PDF bản thảo
        raise SystemExit(f"[LỖI] {e}")

if __name__ == "__main__":
    main()
