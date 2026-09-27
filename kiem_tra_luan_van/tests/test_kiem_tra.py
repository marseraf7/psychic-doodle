# -*- coding: utf-8 -*-
import os, sys
import pytest

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, GOC)
import kiem_tra_luan_van as k

VI_DU = os.path.join(GOC, "vi_du")


def test_tach_cau_khong_cat_sau_viet_tat_va_chu_cai_dau_ten():
    doan = "Л. С. Выготский подчёркивал роль среды, т. е. отношений [3, с. 12]. Второе предложение."
    assert k.tach_cau(doan) == ["Л. С. Выготский подчёркивал роль среды, т. е. отношений [3, с. 12].",
                                "Второе предложение."]


def test_goc_tu_gop_cac_cach():
    assert k.goc_tu("психологии") == k.goc_tu("психология") == k.goc_tu("психологию")


def test_trung_lap_phan_loai_dung():
    kq = k.kiem_tra_trung_lap(os.path.join(VI_DU, "ban_thao.txt"), os.path.join(VI_DU, "nguon"))
    loai = {c["cau"][:25]: c["loai"] for ds in kq["doan"] for c in ds}
    assert loai["Тревожность понимается ка"] == k.THIEU_NGUON
    assert loai["Высокий уровень личностно"] == k.CAN_DIEN_DAT
    assert loai["Как отмечает автор, «учеб"] == k.TRICH_DUNG
    assert loai["Актуальность исследования"] == k.OK
    assert loai["Мы предполагаем, что студ"] == k.OK
    # danh mục tài liệu không bị đem đi so trùng
    assert not any("Список" in c["cau"] for ds in kq["doan"] for c in ds)
    assert any("рaссмотрены" in s for _, s in kq["ky_tu_la"])


def test_trung_lap_nhan_ra_cau_chi_doi_bien_cach(tmp_path):
    (tmp_path / "nguon").mkdir()
    (tmp_path / "nguon" / "a.txt").write_text(
        "Личностная тревожность снижает продуктивность учебной деятельности студентов первого курса.",
        encoding="utf-8")
    (tmp_path / "b.txt").write_text(
        "Личностной тревожностью снижается продуктивность учебных деятельностей студента первых курсов.",
        encoding="utf-8")
    kq = k.kiem_tra_trung_lap(str(tmp_path / "b.txt"), str(tmp_path / "nguon"))
    assert kq["doan"][0][0]["loai"] == k.THIEU_NGUON


def test_bao_cao_html(tmp_path):
    kq = k.kiem_tra_trung_lap(os.path.join(VI_DU, "ban_thao.txt"), os.path.join(VI_DU, "nguon"))
    out = tmp_path / "bc.html"
    k.ghi_bao_cao_html(kq, str(out))
    s = out.read_text(encoding="utf-8")
    assert '<mark class="thieu_nguon"' in s and "istochnik_trevozhnost.txt" in s


def test_kiem_tra_trich_dan():
    kq = k.kiem_tra_trich_dan(os.path.join(VI_DU, "ban_thao.txt"))
    nd = "\n".join(x for _, x in kq["van_de"])
    assert kq["so_muc"] == 4
    assert "[5]" in nd                      # dẫn số không có trong danh mục
    assert "Mục 4 trong danh mục chưa được dẫn" in nd
    assert "gạch nối" in nd                 # mục 1 dùng '-'
    assert "Mục 2 trong danh mục chưa" not in nd


def test_trich_nguyen_van_thieu_nguon_va_so_trang(tmp_path):
    f = tmp_path / "b.txt"
    f.write_text("Автор пишет: «мотивация является ведущим фактором успешности обучения в вузе».\n"
                 "Он добавляет: «внутренняя мотивация важнее внешней для долгосрочных результатов» [1].\n"
                 "Список литературы\nИванов, И. И. Книга. – Москва, 2020. – 100 с.\n", encoding="utf-8")
    nd = [(m, x) for m, x in k.kiem_tra_trich_dan(str(f))["van_de"]]
    assert any(m == "LỖI" and "không có tham chiếu" in x for m, x in nd)
    assert any(m == "NHẮC" and "số trang" in x for m, x in nd)


def test_so_trong_ngoac():
    assert k._so_trong_ngoac("5, с. 12; 7–9") == ({5, 7, 8, 9}, True)
    assert k._so_trong_ngoac("1, 3") == ({1, 3}, False)


@pytest.mark.parametrize("dong, mong_doi", [
    ({"loai": "sach", "tac_gia": "Выготский Лев Семенович", "ten": "Мышление и речь",
      "thanh_pho": "Москва", "nha_xb": "Лабиринт", "nam": "1999", "so_trang": "352"},
     "Выготский, Л. С. Мышление и речь / Л. С. Выготский. – Москва : Лабиринт, 1999. – 352 с."),
    ({"loai": "bai_bao", "tac_gia": "Петрова А.А.", "ten": "Мотивация", "tap_chi": "Вопросы психологии",
      "nam": "2020", "so": "3", "trang": "10-20"},
     "Петрова, А. А. Мотивация / А. А. Петрова // Вопросы психологии. – 2020. – № 3. – С. 10–20."),
    ({"loai": "luan_an", "tac_gia": "Кузнецова Елена Николаевна", "ten": "Тревожность",
      "chuyen_nganh": "5.3.1", "noi_bao_ve": "МГУ", "thanh_pho": "Москва", "nam": "2016", "so_trang": "210"},
     "Кузнецова, Е. Н. Тревожность : специальность 5.3.1 : диссертация на соискание ученой степени "
     "кандидата психологических наук / Кузнецова Елена Николаевна ; МГУ. – Москва, 2016. – 210 с."),
])
def test_dinh_dang_gost(dong, mong_doi):
    assert k.dinh_dang_muc(dong) == mong_doi


def test_dinh_dang_csv_dau_cham_phay(tmp_path):
    f = tmp_path / "t.csv"
    f.write_text("loai;tac_gia;ten;thanh_pho;nam;so_trang\n"
                 "sach;Smith J.;Anxiety;London;2001;90\n"
                 "sach;Орлов О. О.;Стресс;Москва;2010;120\n", encoding="utf-8")
    out = tmp_path / "dm.txt"
    k.main(["dinh-dang", "--csv", str(f), "--out", str(out)])
    dong = out.read_text(encoding="utf-8").splitlines()
    assert dong[2].startswith("1. Орлов")          # tiếng Nga xếp trước
    assert dong[3].endswith("– 90 p.")


def test_docx_chu_trang_va_ky_tu_an(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.shared import RGBColor
    d = docx.Document()
    p = d.add_paragraph("Обычный текст​ исследования. ")
    p.add_run("скрытый текст").font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    path = str(tmp_path / "b.docx")
    d.save(path)
    loi = k.do_chu_an_docx(path) + k.do_ky_tu_la(k.doc_doan_van(path))
    ly_do = " ".join(l for l, _ in loi)
    assert "chữ màu trắng" in ly_do and "zero-width space" in ly_do


# ----------------------------------------------------------------------
# Các trường hợp từ lượt tự review
# ----------------------------------------------------------------------
CAU_1 = ("Тревожность понимается как устойчивая индивидуальная особенность, проявляющаяся в склонности "
         "человека испытывать беспокойство в самых различных жизненных ситуациях.")
CAU_2 = ("Высокий уровень личностной тревожности снижает продуктивность учебной деятельности и затрудняет "
         "адаптацию студентов первого курса к условиям вуза.")
NGUON = os.path.join(VI_DU, "nguon")


def _loai(kq):
    return [c["loai"] for ds in kq["doan"] for c in ds]


def test_tham_chieu_cuoi_doan_ap_cho_ca_doan(tmp_path):
    f = tmp_path / "b.txt"
    f.write_text(f"{CAU_1} {CAU_2[:-1]} [1].\n", encoding="utf-8")
    # đã dẫn nguồn -> không còn đỏ; vẫn vàng vì chép gần nguyên văn
    assert _loai(k.kiem_tra_trung_lap(str(f), NGUON)) == [k.CAN_DIEN_DAT, k.CAN_DIEN_DAT]


def test_tham_chieu_dau_doan_khong_ap_cho_cau_sau(tmp_path):
    f = tmp_path / "b.txt"
    f.write_text(f"Об этом писал Иванов [1]. {CAU_1}\n", encoding="utf-8")
    assert _loai(k.kiem_tra_trung_lap(str(f), NGUON))[-1] == k.THIEU_NGUON


def test_trich_dan_nhieu_cau_trong_ngoac_kep(tmp_path):
    f = tmp_path / "b.txt"
    f.write_text(f"Автор пишет: «{CAU_1} {CAU_2}» [1, с. 5].\n", encoding="utf-8")
    assert set(_loai(k.kiem_tra_trung_lap(str(f), NGUON))) <= {k.TRICH_DUNG, k.OK}


def test_nguon_rong_duoc_canh_bao(tmp_path):
    (tmp_path / "n").mkdir()
    (tmp_path / "n" / "scan.txt").write_text("  \n", encoding="utf-8")
    kq = k.kiem_tra_trung_lap(os.path.join(VI_DU, "ban_thao.txt"), str(tmp_path / "n"))
    assert any("scan.txt" in x and "OCR" in x for x in kq["canh_bao_nguon"])
    out = tmp_path / "bc.html"
    k.ghi_bao_cao_html(kq, str(out))
    assert "không được so sánh" in out.read_text(encoding="utf-8")


def _pdf_chu(path, dong):
    """PDF tối giản có lớp chữ (Helvetica, ASCII) — không cần thư viện tạo PDF."""
    lenh = "BT /F1 11 Tf 40 780 Td 14 TL " + " ".join(f"({d}) '" for d in dong) + " ET"
    obj = ["<< /Type /Catalog /Pages 2 0 R >>",
           "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
           "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
           "/Resources << /Font << /F1 5 0 R >> >> >>",
           f"<< /Length {len(lenh)} >>\nstream\n{lenh}\nendstream",
           "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, vt = b"%PDF-1.4\n", []
    for i, o in enumerate(obj, 1):
        vt.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(obj) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{v:010d} 00000 n \n" for v in vt).encode()
    out += f"trailer\n<< /Size {len(obj) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


def test_nguon_pdf_co_chu_va_pdf_scan(tmp_path):
    pytest.importorskip("pypdf")
    (tmp_path / "n").mkdir()
    _pdf_chu(tmp_path / "n" / "a.pdf", ["Personality anxiety strongly reduces academic performance",
                                         "among first year university students in large cities."])
    _pdf_chu(tmp_path / "n" / "scan.pdf", [])
    f = tmp_path / "b.txt"
    f.write_text("Personality anxiety strongly reduces academic performance among first year "
                 "university students in large cities.\n", encoding="utf-8")
    kq = k.kiem_tra_trung_lap(str(f), str(tmp_path / "n"))
    assert _loai(kq) == [k.THIEU_NGUON] and kq["doan"][0][0]["nguon"] == "a.pdf"
    assert any("scan.pdf" in x for x in kq["canh_bao_nguon"])


def test_txt_ngat_dong_cung():
    doan = k._noi_dong_txt([
        "Введение",
        "Высокий уровень личностной тревожности снижает продуктивность учебной",
        "деятельности студентов первого курса.",
        "Список литературы",
        "Иванов, И. И. Тревожность. – Москва, 2018. – 200 с.",
        "Петров, П. П. Стресс. – Москва, 2015. – 180 с.",
    ])
    assert doan == ["Введение",
                    "Высокий уровень личностной тревожности снижает продуктивность учебной "
                    "деятельности студентов первого курса.",
                    "Список литературы",
                    "Иванов, И. И. Тревожность. – Москва, 2018. – 200 с.",
                    "Петров, П. П. Стресс. – Москва, 2015. – 180 с."]


def test_tach_cau_sau_viet_tat_cuoi_cau():
    assert len(k.tach_cau("Это описано в работах Иванова и др. Далее мы рассмотрим методику.")) == 2
    assert len(k.tach_cau("Данные собраны и т. д. Потом анализ.")) == 2
    assert len(k.tach_cau("См. рис. 3. Далее текст.")) == 2
    assert len(k.tach_cau("Как писал Л. С. Выготский, т. е. автор, это важно.")) == 1


def test_trich_dan_tac_gia_nam(tmp_path):
    f = tmp_path / "b.txt"
    f.write_text("По мнению Выготского (1934), мышление связано с речью.\n"
                 "Мотивация важна (Иванова, 2010; Смирнов и др., 2019).\n"
                 "Список литературы\n"
                 "Выготский, Л. С. Мышление и речь. – Москва, 1934. – 300 с.\n"
                 "Иванова, А. А. Мотивация. – Москва, 2010. – 100 с.\n"
                 "Петров, П. П. Стресс. – Москва, 2011. – 90 с.\n", encoding="utf-8")
    nd = "\n".join(x for _, x in k.kiem_tra_trich_dan(str(f))["van_de"])
    assert "Mục 1 trong danh mục chưa" not in nd and "Mục 2 trong danh mục chưa" not in nd
    assert "Mục 3 trong danh mục chưa được dẫn" in nd
    assert "Смирнов" in nd and "không có tác giả này" in nd


def test_nam_trong_ngoac_vuong_khong_phai_so_thu_tu(tmp_path):
    f = tmp_path / "b.txt"
    f.write_text("Закон принят [2015] и описан [1].\nСписок литературы\n"
                 "Иванов, И. И. Книга. – Москва, 2010. – 100 с.\n", encoding="utf-8")
    assert not [x for m, x in k.kiem_tra_trich_dan(str(f))["van_de"] if m == "LỖI"]


def _docx_co_chu_thich(path, doan, chu_thich):
    """Tạo .docx có chú thích chân trang thật (word/footnotes.xml). Trong `doan`, '{N}' là
    vị trí tham chiếu tới chú thích thứ N (bắt đầu từ 1)."""
    import zipfile, re as _re
    docx = pytest.importorskip("docx")
    d = docx.Document()
    for t in doan:
        d.add_paragraph(t)
    d.save(path)
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(path) as z:
        tep = {n: z.read(n) for n in z.namelist()}
    xml = tep["word/document.xml"].decode("utf-8")
    xml = _re.sub(r"\{(\d+)\}", lambda m: f'</w:t></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/>'
                  f'</w:rPr><w:footnoteReference w:id="{m.group(1)}"/></w:r><w:r><w:t xml:space="preserve">', xml)
    tep["word/document.xml"] = xml.encode("utf-8")
    fn = "".join(f'<w:footnote w:id="{i}"><w:p><w:r><w:t>{t}</w:t></w:r></w:p></w:footnote>'
                 for i, t in enumerate(chu_thich, 1))
    tep["word/footnotes.xml"] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W}">'
                                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                                 f'{fn}</w:footnotes>').encode("utf-8")
    tep["word/_rels/document.xml.rels"] = tep["word/_rels/document.xml.rels"].replace(
        b"</Relationships>", b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/'
        b'2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
    tep["[Content_Types].xml"] = tep["[Content_Types].xml"].replace(
        b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
    with zipfile.ZipFile(path, "w") as z:
        for n, b in tep.items():
            z.writestr(n, b)
    return path


def test_chu_thich_chan_trang(tmp_path):
    p = _docx_co_chu_thich(str(tmp_path / "b.docx"), [
        CAU_1 + "{1}",
        "Выготский писал: «мышление не выражается в слове, но совершается в слове».{2}",
        "Список литературы",
        "Выготский, Л. С. Мышление и речь. – Москва : Лабиринт, 1999. – 352 с.",
        "Иванов, И. И. Тревожность. – Москва, 2018. – 200 с.",
        "Сидоров, П. П. Стресс. – Москва, 2015. – 180 с.",
    ], ["Иванов И. И. Тревожность. М., 2018. С. 12.", "Выготский Л. С. Мышление и речь. М., 1999."])
    assert k.doc_doan_van(p)[0].endswith("⟨1⟩")
    assert k.doc_chu_thich(p)["⟨2⟩"].startswith("Выготский")
    # câu chép có chú thích chân trang -> đã dẫn nguồn (vàng), không phải đỏ
    assert _loai(k.kiem_tra_trung_lap(p, NGUON))[0] == k.CAN_DIEN_DAT
    kq = k.kiem_tra_trich_dan(p)
    nd = "\n".join(x for _, x in kq["van_de"])
    assert kq["so_chu_thich"] == 2
    assert "Mục 3 trong danh mục chưa được dẫn" in nd
    assert "Mục 1 trong danh mục chưa" not in nd and "Mục 2 trong danh mục chưa" not in nd
    assert "nên ghi số trang" in nd and "không có tham chiếu" not in nd   # chú thích 2 thiếu 'С.'


def test_chu_trang_dat_qua_style(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import RGBColor
    d = docx.Document()
    st = d.styles.add_style("Trang", WD_STYLE_TYPE.PARAGRAPH)
    st.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    d.add_paragraph("невидимый текст", style="Trang")
    d.add_paragraph("обычный текст")
    path = str(tmp_path / "b.docx")
    d.save(path)
    assert [s for _, s in k.do_chu_an_docx(path)] == ["невидимый текст"]


def test_ten_viet_chu_tat_truoc_ho():
    assert k.tach_ten("Л. С. Выготский") == ("Выготский", "Л. С.")
    assert k.tach_ten("Л.С. Выготский") == ("Выготский", "Л. С.")


def test_csv_thieu_cot_va_loai_sai_bi_bo_qua(tmp_path, capsys):
    f = tmp_path / "t.csv"
    f.write_bytes("loai;tac_gia;ten;tap_chi;thanh_pho;nam;so_trang;trang\n"
                  "bai_bao;Иванов И. И.;Статья;;;;;\n"
                  "book;Орлов О. О.;Стресс;;Москва;2010;120;\n"
                  "Книга;Петров П. П.;Книга;;Москва;2011;90;\n"
                  "sahc;Петров П. П.;Книга;;Москва;2011;90;\n".encode("cp1251"))   # Excel tiếng Nga
    out = tmp_path / "dm.txt"
    k.main(["dinh-dang", "--csv", str(f), "--out", str(out)])
    ra = capsys.readouterr().out
    assert "Dòng 2" in ra and "tap_chi, nam, trang" in ra
    assert "Dòng 5" in ra and "không hợp lệ" in ra
    muc = out.read_text(encoding="utf-8").splitlines()[2:]
    assert len(muc) == 2 and "//" not in "".join(muc)


def test_chuan_loai():
    assert k.chuan_loai("Sách") == "sach" and k.chuan_loai("bài báo") == "bai_bao"
    assert k.chuan_loai("Статья") == "bai_bao" and k.chuan_loai("luận án") == "luan_an"
    assert k.chuan_loai("sahc") is None and k.chuan_loai("") == ""
